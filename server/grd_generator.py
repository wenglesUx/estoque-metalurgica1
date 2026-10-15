#!/usr/bin/env python3
"""
grd_generator.py
Gerador de GRD usando openpyxl com keep_vba=True.

Recebe um JSON via stdin com os dados da GRD e escreve o arquivo .xlsm em disco,
imprimindo o caminho do arquivo gerado via stdout.

Uso: echo '<json>' | python3 grd_generator.py
"""

import sys
import json
import os
import subprocess
from datetime import datetime

try:
    import openpyxl
except ImportError:
    # Tenta se auto-instalar antes de desistir — cobre o caso do ambiente de
    # deploy não ter rodado (ou ter perdido) o passo de instalação do openpyxl.
    try:
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "--break-system-packages", "--quiet", "openpyxl"]
        )
        import openpyxl
    except Exception as install_err:
        print(json.dumps({
            "error": f"openpyxl não instalado e a auto-instalação falhou: {install_err}. Execute manualmente: pip install openpyxl"
        }), file=sys.stderr)
        sys.exit(1)


def main():
    raw = sys.stdin.read()
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError as e:
        print(json.dumps({"error": f"JSON inválido: {e}"}), file=sys.stderr)
        sys.exit(1)

    template_path = payload["templatePath"]
    output_path = payload["outputPath"]
    numero = payload["numero"]
    data_str = payload["data"]          # ISO string: "2024-02-03"
    cliente = payload["cliente"]
    obra = payload["obra"]
    materiais = payload["materiais"]    # lista de objetos
    hide_data_sheet = bool(payload.get("hideDataSheet", False))
    print_scale = payload.get("printScale")
    fit_to_page = bool(payload.get("fitToPage", False))
    pdf_mode = bool(payload.get("pdfMode", False))

    # Verifica template
    if not os.path.exists(template_path):
        print(json.dumps({"error": f"Template não encontrado: {template_path}"}), file=sys.stderr)
        sys.exit(1)

    # Garante diretório de saída
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    # Abre o template preservando VBA/macros
    wb = openpyxl.load_workbook(template_path, keep_vba=True)

    # ── Aba "Dados" ────────────────────────────────────────────────────────────
    ws_dados = wb["Dados"]

    # Limpa dados anteriores (linha 3 em diante) — mantém linha 1 (título) e 2 (cabeçalhos)
    for row_num in range(3, ws_dados.max_row + 1):
        for col in range(1, 10):  # colunas A-I
            ws_dados.cell(row=row_num, column=col).value = None

    # Preenche com os materiais selecionados
    for idx, mat in enumerate(materiais):
        row_num = idx + 3  # dados começam na linha 3
        ws_dados.cell(row=row_num, column=1).value = mat.get("certificado", "")   # A — Certif N°
        ws_dados.cell(row=row_num, column=2).value = mat.get("tipo", "")           # B — TIPO
        ws_dados.cell(row=row_num, column=3).value = mat.get("bitola", "")         # C — BITOLA
        ws_dados.cell(row=row_num, column=4).value = mat.get("comprimento", "")   # D — COMP (mm)
        ws_dados.cell(row=row_num, column=5).value = mat.get("largura", "")        # E — largura
        ws_dados.cell(row=row_num, column=6).value = mat.get("quantidade", "")    # F — qntd
        ws_dados.cell(row=row_num, column=7).value = mat.get("unidade", "PÇ")     # G — Unid
        ws_dados.cell(row=row_num, column=8).value = mat.get("aco", "")            # H — aço
        ws_dados.cell(row=row_num, column=9).value = mat.get("nf", "")             # I — NF

    # Na cópia usada para gerar o PDF, somente a aba GRD deve ser impressa.
    # O XLSM principal continua com as duas abas visíveis para uso no Excel.
    if hide_data_sheet:
        ws_dados.sheet_state = "hidden"

    # ── Aba "GRD " (nome com espaço) ──────────────────────────────────────────
    # Tenta com e sem espaço no nome
    ws_grd = None
    for name in ["GRD ", "GRD"]:
        if name in wb.sheetnames:
            ws_grd = wb[name]
            break

    if ws_grd is None:
        print(json.dumps({"error": 'Aba "GRD" não encontrada no template.'}), file=sys.stderr)
        sys.exit(1)

    # A exportação PDF usa uma escala própria para ocupar a folha inteira,
    # sem alterar o template XLSM original entregue ao usuário.
    if fit_to_page:
        ws_grd.page_setup.scale = None
        ws_grd.page_setup.fitToWidth = 1
        ws_grd.page_setup.fitToHeight = 1
        ws_grd.sheet_properties.pageSetUpPr.fitToPage = True
    elif print_scale is not None:
        ws_grd.page_setup.scale = int(print_scale)
        ws_grd.sheet_properties.pageSetUpPr.fitToPage = False

    if pdf_mode and materiais:
        # Valores fixos do formulário de referência: via original, uma via,
        # entregue em meio eletrônico.
        ws_grd["G13"] = 0
        ws_grd["H13"] = "O"
        ws_grd["I13"] = 1
        ws_grd["J13"] = "ME"

    # I1 — número do relatório
    ws_grd["I1"] = f"RELATÓRIO: {numero}"

    # J2 — data (como objeto datetime para o Excel reconhecer como data)
    try:
        data_obj = datetime.fromisoformat(data_str)
        ws_grd["J2"] = data_obj
        # Preserva formato de data se a célula já tinha
        if ws_grd["J2"].number_format == "General":
            ws_grd["J2"].number_format = "DD/MM/YYYY"
    except ValueError:
        ws_grd["J2"] = data_str

    # A7 — cliente
    ws_grd["A7"] = f"CLIENTE: {cliente.upper()}"

    # F7 — obra
    ws_grd["F7"] = f"OBRA:  {obra.upper()}"

    # Salva o arquivo — openpyxl com keep_vba preserva o Content_Types correto para .xlsm
    wb.save(output_path)

    # Retorna sucesso via stdout
    print(json.dumps({"success": True, "outputPath": output_path}))


if __name__ == "__main__":
    main()
